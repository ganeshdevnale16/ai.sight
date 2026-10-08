"""Read a sample report (PDF, PowerPoint, Word, Excel/CSV or image) into text + page images
so the model can see its structure, KPIs and charts."""
from __future__ import annotations

import base64
import io
from pathlib import Path

MAX_TEXT = 60000
MAX_IMAGES = 8


def _img_b64(pil, max_side=1500, quality=82):
    pil = pil.convert("RGB")
    w, h = pil.size
    s = min(1.0, max_side / max(w, h))
    if s < 1:
        pil = pil.resize((int(w * s), int(h * s)))
    buf = io.BytesIO()
    pil.save(buf, format="JPEG", quality=quality)
    return {"type": "image", "media_type": "image/jpeg", "data": base64.b64encode(buf.getvalue()).decode()}


def extract(path: Path, log=lambda m: None) -> dict:
    ext = path.suffix.lower()
    text, images = "", []
    if ext == ".pdf":
        import pdfplumber
        import pypdfium2 as pdfium
        chunks = []
        with pdfplumber.open(str(path)) as pdf:
            for i, page in enumerate(pdf.pages[:40], 1):
                chunks.append(f"\n--- Page {i} ---\n{page.extract_text() or ''}")
        text = "".join(chunks)
        doc = pdfium.PdfDocument(str(path))
        for i in range(min(len(doc), MAX_IMAGES)):
            images.append(_img_b64(doc[i].render(scale=1.4).to_pil()))
        log(f"Read {len(chunks)} PDF pages, rendered {len(images)} for visual review")
    elif ext in (".pptx", ".potx"):
        from pptx import Presentation
        prs = Presentation(str(path))
        chunks = []
        for i, slide in enumerate(prs.slides, 1):
            parts = [f"\n--- Slide {i} ---"]
            for sh in slide.shapes:
                if sh.has_text_frame and sh.text_frame.text.strip():
                    parts.append(sh.text_frame.text.strip())
                if getattr(sh, "has_table", False) and sh.has_table:
                    for r in sh.table.rows:
                        parts.append(" | ".join(c.text for c in r.cells))
                if getattr(sh, "has_chart", False) and sh.has_chart:
                    ch = sh.chart
                    title = ch.chart_title.text_frame.text if ch.has_title else ""
                    cats = list(ch.plots[0].categories) if len(ch.plots) else []
                    series = [s.name for p in ch.plots for s in p.series]
                    parts.append(f"[Chart: {ch.chart_type} '{title}' categories={cats[:20]} series={series}]")
                if sh.shape_type == 13 and len(images) < MAX_IMAGES:  # picture
                    try:
                        from PIL import Image
                        images.append(_img_b64(Image.open(io.BytesIO(sh.image.blob))))
                    except Exception:
                        pass
            chunks.append("\n".join(parts))
        text = "".join(chunks)
        log(f"Read {len(prs.slides)} slides")
    elif ext in (".docx", ".dotx"):
        import docx
        d = docx.Document(str(path))
        parts = [p.text for p in d.paragraphs if p.text.strip()]
        for t in d.tables:
            for r in t.rows:
                parts.append(" | ".join(c.text.strip() for c in r.cells))
        text = "\n".join(parts)
        from PIL import Image
        for rel in d.part.rels.values():
            if "image" in rel.reltype and len(images) < MAX_IMAGES:
                try:
                    images.append(_img_b64(Image.open(io.BytesIO(rel.target_part.blob))))
                except Exception:
                    pass
        log(f"Read Word document ({len(parts)} blocks)")
    elif ext in (".xlsx", ".xlsm", ".xls", ".csv"):
        import pandas as pd
        sheets = pd.read_excel(path, sheet_name=None, header=None) if ext != ".csv" else {"Sheet1": pd.read_csv(path, header=None)}
        parts = []
        for name, df in sheets.items():
            parts.append(f"\n--- Sheet {name} ({df.shape[0]} rows x {df.shape[1]} cols) ---")
            parts.append(df.head(60).fillna("").astype(str).to_csv(sep="|", index=False, header=False))
        text = "\n".join(parts)
        log(f"Read {len(sheets)} sheet(s)")
    elif ext in (".png", ".jpg", ".jpeg", ".webp", ".gif"):
        from PIL import Image
        images.append(_img_b64(Image.open(path)))
        log("Read image")
    elif ext in (".txt", ".md"):
        text = path.read_text(encoding="utf-8", errors="replace")
    else:
        raise ValueError("Unsupported sample report type. Use PDF, PPTX, DOCX, XLSX, CSV or an image.")
    return {"text": text[:MAX_TEXT], "images": images}
